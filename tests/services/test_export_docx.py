"""``ExportService`` DOCX 格式集成测试(P6-B-06)。

覆盖维度(对齐 plan §10 测试约定 5 类):

- **happy**:create_export(format=docx) + execute → zip 内含 ``book.docx`` +
  ``meta.json``;``Document()`` 反序列化读取 paragraphs 总数符合 +
  Heading 1 数 = 卷数 + Heading 2 数 = 章节数
- **content**:章节段落 + 标题写入 .docx body
- **volume scope**:scope=volume 章节过滤正确
- **rollback**:render_docx 注入失败 → manifest=failed + 半 zip 删
- **WS event**:4 stage(collecting / rendering / packaging / done)publish 验证

错误码 / scope gate 覆盖沿用 ``test_export_service.py``(本任务不重复),只额外
覆盖 docx 格式专属路径。
"""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable, Iterator
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from src.api import task_registry
from src.core.db import Base
from src.core.db.models import ChapterORM, VolumeORM, WorkspaceORM
from src.core.workspace.workspace_paths import WorkspacePaths
from src.services.export_service import ExportService
from src.services.workspace_service import WorkspaceService


class _FakeWSManager:
    def __init__(self) -> None:
        self.events: list[dict[str, Any]] = []

    def publish(
        self, *, channel: str, event: str, data: dict[str, Any]
    ) -> None:
        self.events.append(
            {"channel": channel, "event": event, "data": data}
        )


@pytest.fixture
def workspace_service(tmp_path: Path) -> WorkspaceService:
    workspaces_root = tmp_path / "workspaces"
    archives_root = tmp_path / "archives"
    workspaces_root.mkdir()
    archives_root.mkdir()
    return WorkspaceService(
        workspaces_root=workspaces_root, archives_root=archives_root
    )


@pytest.fixture
def workspace_id(workspace_service: WorkspaceService) -> str:
    workspace_service.create(
        "ws-docx",
        display_name="DOCX 测试作品",
        genre="xianxia",
        premise="一句话设定",
        activate=False,
    )
    return "ws-docx"


@pytest.fixture
def session_factory() -> Callable[[str], Session]:
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)

    def factory(_workspace_id: str) -> Session:
        return Session(engine)

    return factory


def _seed_db(
    session_factory: Callable[[str], Session],
    workspace_id: str,
    *,
    chapter_count: int = 3,
) -> None:
    session = session_factory(workspace_id)
    try:
        now = datetime.now(timezone.utc)
        session.merge(
            WorkspaceORM(
                workspace_id=workspace_id,
                display_name="DOCX 测试作品",
                genre="xianxia",
                premise="p",
                created_at=now,
                last_active_at=now,
                engine_version="0.1.0",
                status="active",
                total_chapters=chapter_count,
                completed_chapters=0,
                config_overrides={},
            )
        )
        session.merge(
            VolumeORM(
                volume_id="vol-1",
                workspace_id=workspace_id,
                title="开端卷",
                summary="s",
                chapter_count=chapter_count,
                tension_template="heros_journey",
                arcs_json=[],
                tension_curve_json=[],
                created_at=now,
                last_updated_at=now,
            )
        )
        for i in range(1, chapter_count + 1):
            session.merge(
                ChapterORM(
                    chapter_id=f"ch-{i:04d}",
                    workspace_id=workspace_id,
                    current_version=1,
                    status="approved",
                    rewrite_count=0,
                    revise_count=0,
                    created_at=now,
                    approved_at=now,
                    total_token_usage={},
                    total_cost_usd=0.0,
                )
            )
        session.commit()
    finally:
        session.close()


def _write_final(
    workspace_service: WorkspaceService,
    workspace_id: str,
    chapter_id: str,
    content: str,
) -> None:
    paths = WorkspacePaths(
        workspace_id, workspaces_root=workspace_service.workspaces_root
    )
    paths.chapter_dir(chapter_id).mkdir(parents=True, exist_ok=True)
    (paths.chapter_dir(chapter_id) / "final.md").write_text(
        content, encoding="utf-8"
    )


def _write_card(
    workspace_service: WorkspaceService,
    workspace_id: str,
    chapter_id: str,
    *,
    title: str,
    chapter_number: int,
) -> None:
    paths = WorkspacePaths(
        workspace_id, workspaces_root=workspace_service.workspaces_root
    )
    paths.cards_dir.mkdir(parents=True, exist_ok=True)
    (paths.cards_dir / f"{chapter_id}.json").write_text(
        json.dumps(
            {
                "chapter_id": chapter_id,
                "title": title,
                "chapter_number": chapter_number,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


@pytest.fixture
def export_service(
    workspace_service: WorkspaceService,
    session_factory: Callable[[str], Session],
) -> ExportService:
    return ExportService(
        workspace_service=workspace_service,
        session_factory=session_factory,
        ws_manager=None,
        engine_version="0.1.0",
    )


@pytest.fixture(autouse=True)
def clear_task_registry() -> Iterator[None]:
    task_registry.clear()
    yield
    task_registry.clear()


def _seed_three_chapters(
    workspace_service: WorkspaceService,
    workspace_id: str,
) -> None:
    """3 章正文 + 卡片(共享工具,2 段正文 + 1 空 final)。"""
    _write_final(
        workspace_service, workspace_id, "ch-0001", "正文段一\n\n正文段二"
    )
    _write_final(
        workspace_service, workspace_id, "ch-0002", "另一章\n\n第二段"
    )
    _write_final(workspace_service, workspace_id, "ch-0003", "")  # 空 final
    _write_card(
        workspace_service,
        workspace_id,
        "ch-0001",
        title="启程",
        chapter_number=1,
    )
    _write_card(
        workspace_service,
        workspace_id,
        "ch-0002",
        title="风波",
        chapter_number=2,
    )
    _write_card(
        workspace_service,
        workspace_id,
        "ch-0003",
        title="转折",
        chapter_number=3,
    )


def _read_docx_from_zip(zip_bytes: bytes) -> Document:
    """zip → 解 ``book.docx`` → ``Document()`` 反序列化。"""
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        docx_bytes = zf.read("book.docx")
    return Document(io.BytesIO(docx_bytes))


# ---- happy:format=docx scope=full → round-trip ----------------------------


class TestExecuteDocx:
    def test_docx_full_scope_writes_zip_with_book_and_meta(
        self,
        export_service: ExportService,
        workspace_service: WorkspaceService,
        session_factory: Callable[[str], Session],
        workspace_id: str,
    ) -> None:
        """zip 内含 ``book.docx`` + ``meta.json``;Heading 1 = 卷数(1)+
        Heading 2 = 章节数(3);paragraphs 总数 = TOC 占位 + 2+2+1 段 = 6。"""
        _seed_db(session_factory, workspace_id, chapter_count=3)
        _seed_three_chapters(workspace_service, workspace_id)

        task_id = task_registry.new_task_id()
        task_registry.record_pending(task_id, type="workspace_export")
        export_id, manifest = export_service.create_export(
            workspace_id,
            task_id=task_id,
            format="docx",
            scope="full",
            volume_id=None,
            start_chapter_id=None,
            end_chapter_id=None,
            include_metadata=True,
            include_review_history=False,
        )
        assert export_id.startswith("exp_")
        assert manifest["status"] == "pending"

        result = export_service.execute(
            workspace_id, task_id=task_id, export_id=export_id
        )
        assert result["status"] == "done"
        assert result["chapter_count"] == 3
        assert result["format"] == "docx"

        zip_path = export_service.get_download_path(workspace_id, export_id)
        zip_bytes = zip_path.read_bytes()
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            names = zf.namelist()
            assert "book.docx" in names
            assert "meta.json" in names
            meta = json.loads(zf.read("meta.json").decode("utf-8"))
            assert meta["format"] == "docx"
            assert meta["chapter_count"] == 3
            assert meta["engine_version"] == "0.1.0"

        # 反序列化 + Heading / 段落计数验收
        doc = _read_docx_from_zip(zip_bytes)
        heading_1_count = sum(
            1 for p in doc.paragraphs if p.style.name == "Heading 1"
        )
        heading_2_count = sum(
            1 for p in doc.paragraphs if p.style.name == "Heading 2"
        )
        assert heading_1_count == 1  # 1 卷
        assert heading_2_count == 3  # 3 章
        # Normal paragraphs:TOC 占位 + ch-0001 (2 段) + ch-0002 (2 段) +
        # ch-0003 (1 段空占位) = 6
        normal_count = sum(
            1 for p in doc.paragraphs if p.style.name == "Normal"
        )
        assert normal_count == 6

        # task_registry done
        info = task_registry.get(task_id)
        assert info is not None
        assert info.status == "done"

    def test_docx_chapter_content_contains_title_and_paragraphs(
        self,
        export_service: ExportService,
        workspace_service: WorkspaceService,
        session_factory: Callable[[str], Session],
        workspace_id: str,
    ) -> None:
        """章节 .docx body 含 ``第 N 章: 标题`` Heading 2 + 段落原文。"""
        _seed_db(session_factory, workspace_id, chapter_count=1)
        _write_final(
            workspace_service,
            workspace_id,
            "ch-0001",
            "第一段落\n\n第二段落",
        )
        _write_card(
            workspace_service,
            workspace_id,
            "ch-0001",
            title="启程",
            chapter_number=1,
        )

        task_id = task_registry.new_task_id()
        task_registry.record_pending(task_id, type="workspace_export")
        export_id, _ = export_service.create_export(
            workspace_id,
            task_id=task_id,
            format="docx",
            scope="full",
            volume_id=None,
            start_chapter_id=None,
            end_chapter_id=None,
            include_metadata=True,
            include_review_history=False,
        )
        export_service.execute(
            workspace_id, task_id=task_id, export_id=export_id
        )
        zip_path = export_service.get_download_path(workspace_id, export_id)
        doc = _read_docx_from_zip(zip_path.read_bytes())
        body_texts = [p.text for p in doc.paragraphs]
        assert "第 1 章: 启程" in body_texts
        assert "第一段落" in body_texts
        assert "第二段落" in body_texts

    def test_docx_includes_toc_placeholder_paragraph(
        self,
        export_service: ExportService,
        workspace_service: WorkspaceService,
        session_factory: Callable[[str], Session],
        workspace_id: str,
    ) -> None:
        """TOC 占位段落字面"[请在 Word 中按 F9 更新目录]"出现在 .docx 中
        (自决项 8;python-docx 不原生支持 TOC field)。"""
        _seed_db(session_factory, workspace_id, chapter_count=1)
        _write_final(workspace_service, workspace_id, "ch-0001", "x")
        _write_card(
            workspace_service,
            workspace_id,
            "ch-0001",
            title="启程",
            chapter_number=1,
        )

        task_id = task_registry.new_task_id()
        task_registry.record_pending(task_id, type="workspace_export")
        export_id, _ = export_service.create_export(
            workspace_id,
            task_id=task_id,
            format="docx",
            scope="full",
            volume_id=None,
            start_chapter_id=None,
            end_chapter_id=None,
            include_metadata=True,
            include_review_history=False,
        )
        export_service.execute(
            workspace_id, task_id=task_id, export_id=export_id
        )
        zip_path = export_service.get_download_path(workspace_id, export_id)
        doc = _read_docx_from_zip(zip_path.read_bytes())
        body_texts = [p.text for p in doc.paragraphs]
        assert "[请在 Word 中按 F9 更新目录]" in body_texts

    def test_docx_volume_scope_excludes_other_volumes(
        self,
        export_service: ExportService,
        workspace_service: WorkspaceService,
        session_factory: Callable[[str], Session],
        workspace_id: str,
    ) -> None:
        _seed_db(session_factory, workspace_id, chapter_count=2)
        _seed_three_chapters(workspace_service, workspace_id)
        task_id = task_registry.new_task_id()
        task_registry.record_pending(task_id, type="workspace_export")
        export_id, _ = export_service.create_export(
            workspace_id,
            task_id=task_id,
            format="docx",
            scope="volume",
            volume_id="vol-1",
            start_chapter_id=None,
            end_chapter_id=None,
            include_metadata=True,
            include_review_history=False,
        )
        result = export_service.execute(
            workspace_id, task_id=task_id, export_id=export_id
        )
        # chapter_count 取决于 vol-1 内章节 = 2(seed_db chapter_count=2)
        assert result["chapter_count"] == 2


# ---- rollback:render_docx 注入失败 ----------------------------------------


class TestRollback:
    def test_docx_execute_failure_marks_manifest_failed_no_zip(
        self,
        workspace_service: WorkspaceService,
        session_factory: Callable[[str], Session],
        workspace_id: str,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        _seed_db(session_factory, workspace_id, chapter_count=1)
        _write_final(workspace_service, workspace_id, "ch-0001", "x")
        service = ExportService(
            workspace_service=workspace_service,
            session_factory=session_factory,
            ws_manager=None,
        )
        task_id = task_registry.new_task_id()
        task_registry.record_pending(task_id, type="workspace_export")
        export_id, _ = service.create_export(
            workspace_id,
            task_id=task_id,
            format="docx",
            scope="full",
            volume_id=None,
            start_chapter_id=None,
            end_chapter_id=None,
            include_metadata=True,
            include_review_history=False,
        )

        # 注入失败:monkeypatch render_docx 抛错
        from src.services import export_service as svc_mod

        def boom(*args: Any, **kwargs: Any) -> None:
            raise RuntimeError("simulated docx render failure")

        monkeypatch.setattr(svc_mod, "render_docx", boom)

        result = service.execute(
            workspace_id, task_id=task_id, export_id=export_id
        )
        assert result["status"] == "failed"

        paths = WorkspacePaths(
            workspace_id, workspaces_root=workspace_service.workspaces_root
        )
        zip_path = paths.exports_dir / export_id / "export.zip"
        assert not zip_path.exists()
        manifest_path = paths.exports_dir / export_id / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        assert manifest["status"] == "failed"
        assert "simulated docx render failure" in manifest["error"]

        info = task_registry.get(task_id)
        assert info is not None
        assert info.status == "failed"


# ---- WS event:4 stage(同 md / epub 模式) --------------------------------


class TestWsEventsDocx:
    def test_publish_four_stages_in_order_docx(
        self,
        workspace_service: WorkspaceService,
        session_factory: Callable[[str], Session],
        workspace_id: str,
    ) -> None:
        _seed_db(session_factory, workspace_id, chapter_count=1)
        _write_final(workspace_service, workspace_id, "ch-0001", "x")
        _write_card(
            workspace_service,
            workspace_id,
            "ch-0001",
            title="启程",
            chapter_number=1,
        )
        fake_ws = _FakeWSManager()
        service = ExportService(
            workspace_service=workspace_service,
            session_factory=session_factory,
            ws_manager=fake_ws,  # type: ignore[arg-type]
        )
        task_id = task_registry.new_task_id()
        task_registry.record_pending(task_id, type="workspace_export")
        export_id, _ = service.create_export(
            workspace_id,
            task_id=task_id,
            format="docx",
            scope="full",
            volume_id=None,
            start_chapter_id=None,
            end_chapter_id=None,
            include_metadata=True,
            include_review_history=False,
        )
        service.execute(
            workspace_id, task_id=task_id, export_id=export_id
        )
        stages = [e["data"]["stage"] for e in fake_ws.events]
        assert stages == ["collecting", "rendering", "packaging", "done"]
        assert all(
            e["channel"] == f"task.{task_id}.progress" for e in fake_ws.events
        )
